# app/services/pdf/invoice_report.py
"""Renders the admin invoice report (rows from fee_service.get_invoice_report_rows) as XLSX or PDF."""
from datetime import datetime
from io import BytesIO
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

MONEY_COLUMNS = ("Total (NGN)", "Paid (NGN)", "Balance (NGN)")


def render_xlsx(rows: List[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    headers = list(rows[0].keys()) if rows else ["No invoices match this filter"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row[h] for h in headers])
    for col, header in enumerate(headers, start=1):
        if header in MONEY_COLUMNS:
            for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                cell.number_format = "#,##0.00"
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = max(12, len(header) + 4)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def render_pdf(rows: List[dict], title: str) -> bytes:
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=0.4 * inch, rightMargin=0.4 * inch,
                            topMargin=0.5 * inch, bottomMargin=0.5 * inch)
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]),
                Paragraph(f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} · {len(rows)} invoices", styles["Normal"]),
                Spacer(1, 0.2 * inch)]
    if rows:
        headers = list(rows[0].keys())
        data = [headers] + [[f"{r[h]:,.2f}" if h in MONEY_COLUMNS else str(r[h]) for h in headers] for r in rows]
        totals = ["Total", "", "", "", ""] + [f"{sum(r[h] for r in rows):,.2f}" for h in MONEY_COLUMNS] + ["", ""]
        data.append(totals)
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#000080")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("ALIGN", (5, 1), (7, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F2F2F2")]),
        ]))
        elements.append(table)
    else:
        elements.append(Paragraph("No invoices match this filter.", styles["Normal"]))
    doc.build(elements)
    return out.getvalue()
